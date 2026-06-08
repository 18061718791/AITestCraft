#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown to Word Converter
将Markdown文件转换为Word文档，字体统一为微软雅黑
优化流程图展示效果
"""

import sys
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_font(cell, font_name='微软雅黑', font_size=10.5, bold=False, color=None):
    """设置单元格字体"""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.bold = bold
            if color:
                run.font.color.rgb = color
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def set_run_font(run, font_name='微软雅黑', font_size=10.5, bold=False, color=None):
    """设置run字体"""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def add_flowchart_table(doc, title, steps):
    """
    添加流程图样式的表格
    steps: list of dict, each dict contains 'name' and 'details'
    """
    # 添加标题
    p = doc.add_paragraph()
    run = p.add_run(title)
    set_run_font(run, font_size=12, bold=True, color=RGBColor(0, 112, 192))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    
    # 创建表格 - 横向流程
    table = doc.add_table(rows=2, cols=len(steps))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 设置表格样式
    table.style = 'Table Grid'
    
    # 第一行：步骤名称（带背景色）
    for i, step in enumerate(steps):
        cell = table.cell(0, i)
        cell.text = step['name']
        set_cell_shading(cell, '4472C4')  # 蓝色背景
        set_cell_font(cell, font_size=11, bold=True, color=RGBColor(255, 255, 255))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 设置单元格宽度
        cell.width = Inches(1.5)
    
    # 第二行：详细说明
    for i, step in enumerate(steps):
        cell = table.cell(1, i)
        cell.text = step['details']
        set_cell_shading(cell, 'D9E2F3')  # 浅蓝色背景
        set_cell_font(cell, font_size=9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # 添加空行

def add_vertical_flowchart(doc, title, steps):
    """
    添加纵向流程图
    steps: list of dict, each dict contains 'name', 'details', 'method'
    """
    # 添加标题
    p = doc.add_paragraph()
    run = p.add_run(title)
    set_run_font(run, font_size=12, bold=True, color=RGBColor(0, 112, 192))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    
    # 创建表格 - 纵向流程
    table = doc.add_table(rows=len(steps) * 2 - 1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.style = 'Table Grid'
    
    row_idx = 0
    for i, step in enumerate(steps):
        # 步骤名称行
        cell = table.cell(row_idx, 0)
        cell.text = f"步骤 {i+1}: {step['name']}"
        set_cell_shading(cell, '4472C4')
        set_cell_font(cell, font_size=11, bold=True, color=RGBColor(255, 255, 255))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.width = Inches(5)
        row_idx += 1
        
        # 详细说明行
        if i < len(steps) - 1:
            cell = table.cell(row_idx, 0)
            details = step.get('details', '')
            if step.get('method'):
                details += f"\n设计方法: {step['method']}"
            cell.text = details
            set_cell_shading(cell, 'D9E2F3')
            set_cell_font(cell, font_size=9)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            row_idx += 1
    
    doc.add_paragraph()

def add_exception_table(doc, exceptions):
    """添加异常处理表格"""
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 表头
    headers = ['异常场景', '处理方式', '用户感知']
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_shading(hdr_cells[i], '70AD47')  # 绿色背景
        set_cell_font(hdr_cells[i], font_size=10, bold=True, color=RGBColor(255, 255, 255))
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 数据行
    for exc in exceptions:
        row_cells = table.add_row().cells
        row_cells[0].text = exc['scene']
        row_cells[1].text = exc['handle']
        row_cells[2].text = exc['feel']
        
        for cell in row_cells:
            set_cell_font(cell, font_size=9)
            set_cell_shading(cell, 'E2EFDA')  # 浅绿色背景
    
    doc.add_paragraph()

def parse_markdown(md_content):
    """解析Markdown内容"""
    lines = md_content.split('\n')
    elements = []
    i = 0
    
    while i < len(lines):
        line = lines[i]
        
        # 跳过空行
        if not line.strip():
            i += 1
            continue
        
        # 标题
        if line.startswith('# '):
            elements.append(('h1', line[2:].strip()))
        elif line.startswith('## '):
            elements.append(('h2', line[3:].strip()))
        elif line.startswith('### '):
            elements.append(('h3', line[4:].strip()))
        elif line.startswith('#### '):
            elements.append(('h4', line[5:].strip()))
        
        # 代码块
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            elements.append(('code', '\n'.join(code_lines)))
        
        # 表格
        elif line.startswith('|') and i + 1 < len(lines) and '---' in lines[i + 1]:
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1
            elements.append(('table', table_lines))
            continue
        
        # 列表项
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            elements.append(('bullet', line.strip()[2:]))
        elif re.match(r'^\s*\d+\.\s', line):
            match = re.match(r'^\s*(\d+)\.\s(.*)', line)
            if match:
                elements.append(('numbered', match.group(2)))
        
        # 普通段落
        else:
            text = line.strip()
            if text:
                elements.append(('paragraph', text))
        
        i += 1
    
    return elements

def create_word_doc(elements, output_path):
    """创建Word文档"""
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    
    # 添加封面标题
    title_para = doc.add_paragraph()
    title_run = title_para.add_run('AI智能体场景分析：测试用例自动生成')
    set_run_font(title_run, font_size=22, bold=True, color=RGBColor(0, 112, 192))
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(30)
    
    # 添加分隔线
    doc.add_paragraph('_' * 50)
    
    for elem_type, content in elements:
        if elem_type == 'h1':
            p = doc.add_heading(content, level=1)
            for run in p.runs:
                set_run_font(run, font_size=18, bold=True, color=RGBColor(0, 112, 192))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(12)
            
        elif elem_type == 'h2':
            p = doc.add_heading(content, level=2)
            for run in p.runs:
                set_run_font(run, font_size=16, bold=True, color=RGBColor(0, 112, 192))
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            
        elif elem_type == 'h3':
            p = doc.add_heading(content, level=3)
            for run in p.runs:
                set_run_font(run, font_size=14, bold=True, color=RGBColor(0, 112, 192))
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            
        elif elem_type == 'h4':
            p = doc.add_heading(content, level=4)
            for run in p.runs:
                set_run_font(run, font_size=12, bold=True, color=RGBColor(0, 112, 192))
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            
        elif elem_type == 'code':
            p = doc.add_paragraph()
            run = p.add_run(content)
            set_run_font(run, font_size=9, color=RGBColor(50, 50, 50))
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)
            # 添加灰色背景
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F5F5F5')
            p.paragraph_format.element.get_or_add_pPr().append(shading)
            
        elif elem_type == 'table':
            table_lines = content
            if len(table_lines) >= 2:
                headers = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = 'Light Grid Accent 1'
                table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                hdr_cells = table.rows[0].cells
                for i, header in enumerate(headers):
                    hdr_cells[i].text = header
                    set_cell_shading(hdr_cells[i], '4472C4')
                    set_cell_font(hdr_cells[i], font_size=10, bold=True, color=RGBColor(255, 255, 255))
                    hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                for line in table_lines[2:]:
                    cells = [cell.strip() for cell in line.split('|')[1:-1]]
                    if cells and any(cells):
                        row_cells = table.add_row().cells
                        for i, cell_text in enumerate(cells):
                            if i < len(row_cells):
                                row_cells[i].text = cell_text
                                set_cell_font(row_cells[i], font_size=9)
                                set_cell_shading(row_cells[i], 'D9E2F3')
                
                doc.add_paragraph()
                
        elif elem_type == 'bullet':
            p = doc.add_paragraph(content, style='List Bullet')
            for run in p.runs:
                set_run_font(run)
            
        elif elem_type == 'numbered':
            p = doc.add_paragraph(content, style='List Number')
            for run in p.runs:
                set_run_font(run)
            
        elif elem_type == 'paragraph':
            p = doc.add_paragraph()
            text = content
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    set_run_font(run, bold=True)
                else:
                    run = p.add_run(part)
                    set_run_font(run)
    
    # 添加专业的流程图
    doc.add_page_break()
    
    # 1. 整体流程架构
    add_flowchart_table(doc, '整体流程架构', [
        {'name': '触发条件', 'details': '用户主动触发\n选择系统/模块\n上传需求文档'},
        {'name': '输入处理', 'details': '文档解析\n需求理解\n场景识别'},
        {'name': '核心处理', 'details': '测试点生成\n测试用例生成\n设计方法识别'},
        {'name': '输出交付', 'details': '用例确认\n导出/保存\nExcel输出'}
    ])
    
    # 2. 核心处理流程
    add_vertical_flowchart(doc, '核心处理流程', [
        {'name': '需求理解', 'details': 'LLM分析需求，提取功能点、业务规则、数据约束', 'method': 'NLP分析'},
        {'name': '测试点生成', 'details': 'LLM基于测试设计方法生成测试点', 'method': '等价类/边界值/场景法'},
        {'name': '设计方法识别', 'details': '基于关键词匹配自动识别设计方法', 'method': '规则引擎'},
        {'name': '用户确认', 'details': '展示测试点，等待用户选择/编辑/确认', 'method': '人工审核'},
        {'name': '测试用例生成', 'details': 'LLM为每个测试点生成完整测试用例', 'method': 'LLM生成'},
        {'name': '格式化输出', 'details': '统一编号、格式校验、生成Excel', 'method': '自动处理'}
    ])
    
    # 3. Dify工作流节点
    add_flowchart_table(doc, 'Dify工作流节点设计', [
        {'name': '开始节点', 'details': '接收用户输入\n需求文本/文档'},
        {'name': '文档解析', 'details': 'Python代码节点\n解析PDF/Word'},
        {'name': '需求分析', 'details': 'LLM节点\n提取功能点'},
        {'name': '测试点生成', 'details': 'LLM节点\n生成测试点'},
        {'name': '方法识别', 'details': '代码节点\n关键词匹配'},
        {'name': '用例生成', 'details': 'LLM节点\n生成用例'},
        {'name': '格式化', 'details': '代码节点\n统一格式'},
        {'name': '导出', 'details': '代码节点\n生成Excel'}
    ])
    
    # 4. 异常处理
    add_exception_table(doc, [
        {'scene': 'LLM API超时', 'handle': '重试3次，间隔2秒；切换备用Provider', 'feel': '显示"正在尝试其他模型..."'},
        {'scene': 'LLM返回格式错误', 'handle': '正则提取JSON；失败则文本解析', 'feel': '无感知，后台自动处理'},
        {'scene': '生成测试点过少(<3条)', 'handle': '自动补充提示词，扩展场景', 'feel': '显示"正在补充测试场景..."'},
        {'scene': '生成测试点过多(>50条)', 'handle': '按模块分批处理，每批10-15条', 'feel': '显示进度条，分批加载'},
        {'scene': '文档解析失败', 'handle': '回退到原始文本，跳过章节提取', 'feel': '提示"已使用原始文本模式"'},
        {'scene': '用户未选择测试点', 'handle': '禁用"生成用例"按钮', 'feel': '按钮置灰+Tooltip提示'},
        {'scene': '用例编号重复', 'handle': '后端重新计算编号', 'feel': '无感知，自动修正'}
    ])
    
    # 保存文档
    doc.save(output_path)
    print(f"Word文档已生成: {output_path}")

def main():
    md_file = r'd:\自动化测试平台\AITestCraft-Tech-style\docs\AI智能体场景分析-测试用例自动生成.md'
    output_file = r'd:\自动化测试平台\AITestCraft-Tech-style\docs\AI智能体场景分析-测试用例自动生成.docx'
    
    with open(md_file, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    elements = parse_markdown(md_content)
    create_word_doc(elements, output_file)

if __name__ == '__main__':
    main()
